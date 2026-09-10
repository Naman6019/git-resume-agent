import os
import pytest
from git_resume.config import load_config
from git_resume.agents.inspector import InspectorAgent
from git_resume.agents.verifier import GroundingVerifierAgent
from git_resume.agents.synthesizer import SynthesizerAgent
from git_resume.agents.schema_discoverer import SchemaDiscoverer

def test_load_config():
    config = load_config('gitresume.yaml')
    assert config.developer.name == 'Naman Manocha'
    assert len(config.repositories) >= 2
    assert len(config.personas) >= 3
    
    # Check link properties on repositories
    funders = next(r for r in config.repositories if r.name == 'FundersAI')
    assert funders.is_deployed is True
    assert funders.live_url == 'https://www.fundersai.co.in'
    assert funders.repo_url == 'https://github.com/Naman6019/FundersAI'
    assert "Live: https://www.fundersai.co.in" in funders.formatted_links
    assert "Repo: https://github.com/Naman6019/FundersAI" in funders.formatted_links

    careflow = next((r for r in config.repositories if r.name == 'CareFlow'), None)
    if careflow:
        assert careflow.is_deployed is False
        assert careflow.repo_url == 'https://github.com/Naman6019/CareFlow-Intelligence'
        assert "Repo: https://github.com/Naman6019/CareFlow-Intelligence" in careflow.formatted_links
        assert "Live:" not in careflow.formatted_links

def test_inspector_agent():
    inspector = InspectorAgent()
    config = load_config('gitresume.yaml')
    stats = inspector.inspect_all(config.repositories)
    assert 'FundersAI' in stats
    assert stats['FundersAI']['commits'] > 0
    assert stats['FundersAI']['loc'] > 0
    assert 'repo_url' in stats['FundersAI']
    assert 'deployed' in stats['FundersAI']
    assert 'formatted_links' in stats['FundersAI']

def test_grounding_verifier():
    verifier = GroundingVerifierAgent()
    ground_truth = {'loc_k': '130K', 'stack': ['Python', 'FastAPI']}
    valid, msg = verifier.verify_bullet("Engineered system with 130K lines of code.", ground_truth)
    assert valid is True
    invalid, msg = verifier.verify_bullet("Engineered system with 900K lines of code.", ground_truth)
    assert invalid is False

def test_synthesizer_agent():
    synthesizer = SynthesizerAgent()
    bullet = synthesizer.synthesize_fundersai({'loc_k': '145K', 'files': 760, 'commits': 402, 'test_suites': 120})
    assert '145K lines of code' in bullet
    assert 'OpenAI Build Week' in bullet

def test_schema_discoverer():
    discoverer = SchemaDiscoverer()
    funders_meta = discoverer.discover_repo_metadata(r"C:\Users\naman\OneDrive\Desktop\FundersAI", "FundersAI")
    assert "FastAPI" in funders_meta["primary_stack"]
    assert "React" in funders_meta["primary_stack"] or "Python" in funders_meta["primary_stack"]
    assert funders_meta["tag"] == "OpenAI Build Week Submission"
    assert funders_meta["live_url"] == "https://www.fundersai.co.in"
    assert funders_meta["repo_url"] == "https://github.com/Naman6019/FundersAI"
    assert funders_meta["deployed"] is True

def test_find_config_path_and_git_utils():
    from git_resume.config import find_config_path
    from git_resume.utils.git_utils import is_git_repo, get_git_remote_details, set_git_config, get_git_config
    
    # 1. Config path resolution
    resolved = find_config_path()
    assert os.path.exists(resolved)
    assert resolved.endswith("gitresume.yaml")

    # 2. Git repo check on known repos
    funders_path = r"C:\Users\naman\OneDrive\Desktop\FundersAI"
    assert is_git_repo(funders_path) is True
    assert is_git_repo(r"C:\NonExistentDirectory_12345") is False

    # 3. Remote details inspection
    details = get_git_remote_details(funders_path)
    assert details["is_git"] is True
    assert details["has_remote"] is True
    assert details["is_github"] is True
    assert "FundersAI" in (details["repo_identifier"] or "")

    # 4. Git config get/set
    resume_path = r"C:\Users\naman\OneDrive\Desktop\resume_automation"
    if is_git_repo(resume_path):
        assert set_git_config(resume_path, "gitresume.testkey", "testvalue") is True
        val = get_git_config("gitresume.testkey", resume_path)
        assert val == "testvalue"

def test_master_personas_config_and_compiler():
    from git_resume.compilers.docx_compiler import DocxCompiler
    config = load_config('gitresume.yaml')
    
    # Check email
    assert config.developer.email == 'namanmanocha42248@gmail.com'
    
    # Check personas
    persona_ids = {p.id: p for p in config.personas}
    assert 'master' in persona_ids
    assert 'master_1page' in persona_ids
    assert persona_ids['master'].resume_file == 'Naman_Manocha_Master_Resume.docx'
    assert persona_ids['master_1page'].resume_file == 'Naman_Manocha_Master_1Page_Resume.docx'
    
    # Check that both files exist in resume_dir
    p2_path = os.path.join(config.output.resume_dir, persona_ids['master'].resume_file)
    p1_path = os.path.join(config.output.resume_dir, persona_ids['master_1page'].resume_file)
    assert os.path.exists(p2_path), f"2-page master not found at {p2_path}"
    assert os.path.exists(p1_path), f"1-page master not found at {p1_path}"
    
    # Test compiler on both
    compiler = DocxCompiler()
    mock_stats = {
        "FundersAI": {"loc_k": "145K", "files": 811, "commits": 416, "test_suites": 147, "deployed": True, "live_url": "https://www.fundersai.co.in", "repo_url": "https://github.com/Naman6019/FundersAI"},
        "TalentOS": {"loc_k": "32K", "files": 134, "commits": 51, "test_suites": 235, "deployed": True, "live_url": "https://all-things-agentic--allthingsagentic-505213.asia-southeast1.hosted.app", "repo_url": "https://github.com/Naman6019/all-things-agentic"}
    }
    assert compiler.update_resume(p2_path, "master", mock_stats) is True
    assert compiler.update_resume(p1_path, "master_1page", mock_stats) is True

def test_scoped_update_resume():
    """target_repo should touch only that project's section, and be a no-op
    (no save, returns False) for a persona that doesn't mention it."""
    import shutil
    import docx
    from git_resume.compilers.docx_compiler import DocxCompiler

    config = load_config('gitresume.yaml')
    persona = next(p for p in config.personas if p.id == 'fde')
    src = os.path.join(config.output.resume_dir, persona.resume_file)
    scratch = os.path.join(os.path.dirname(os.path.abspath(__file__)), f"_scoped_scratch_{persona.resume_file}")
    shutil.copy2(src, scratch)

    try:
        before_texts = [p.text for p in docx.Document(scratch).paragraphs]

        compiler = DocxCompiler()
        mock_stats = {
            "FundersAI": {"loc_k": "999K", "files": 1, "commits": 1, "test_suites": 1},
            "TalentOS": {"loc_k": "32K", "files": 134, "commits": 51},
        }

        # A repo this persona never mentions -> nothing written, file untouched.
        assert compiler.update_resume(scratch, "fde", mock_stats, target_repo="CareFlow") is False
        assert [p.text for p in docx.Document(scratch).paragraphs] == before_texts

        # Scoping to FundersAI updates only FundersAI's paragraph(s).
        assert compiler.update_resume(scratch, "fde", mock_stats, target_repo="FundersAI") is True
        after_texts = [p.text for p in docx.Document(scratch).paragraphs]
        assert any("999K" in t for t in after_texts)
        assert [t for t in after_texts if "TalentOS" in t] == [t for t in before_texts if "TalentOS" in t]
    finally:
        os.remove(scratch)

def test_init_command():
    from typer.testing import CliRunner
    from git_resume.cli import app
    runner = CliRunner()
    
    target_yaml = os.path.join(os.path.dirname(os.path.abspath(__file__)), "_scratch_init.yaml")
    if os.path.exists(target_yaml):
        os.remove(target_yaml)
        
    try:
        result = runner.invoke(app, ["init", "--output", target_yaml])
        assert result.exit_code == 0
        assert os.path.exists(target_yaml)
        
        # Verify the generated starter YAML is valid and parseable by load_config
        cfg = load_config(target_yaml)
        assert cfg.developer.name == "Your Name"
        assert len(cfg.repositories) == 1
        assert cfg.repositories[0].name == "MyProject"
        assert cfg.output.resume_dir == "./resumes"
    finally:
        if os.path.exists(target_yaml):
            os.remove(target_yaml)



